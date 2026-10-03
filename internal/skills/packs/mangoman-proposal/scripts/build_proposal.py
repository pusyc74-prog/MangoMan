"""Build a branded business proposal (PDF, Word and HTML) from proposal.json.

Usage: python3 build_proposal.py proposal.json --out proposal [--no-pdf] [--no-docx]
Writes <out>.pdf (A4, cover plus numbered pages), <out>.docx (editable Word),
<out>.html and <out>.numbers.json (every computed figure, for the checker).

The model writes the words and the price lines (quantity and rate); this
script does all arithmetic: line amounts, discounts, tax (GST by default),
totals, payment milestones and the project duration. So no total can be
typed wrong.
"""
import datetime
import html
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import brandkit as BK  # noqa: E402
import render  # noqa: E402
import vizlib as V  # noqa: E402

SECTION_TYPES = ("summary", "understanding", "scope", "timeline", "pricing", "team", "proof", "terms", "text", "acceptance")
WEB_TYPES = {
    "modern": ('"Segoe UI", Calibri, Carlito, Roboto, Arial, sans-serif', '"Segoe UI", Calibri, Carlito, Roboto, Arial, sans-serif', "Calibri", "Calibri"),
    "editorial": ('Georgia, Cambria, Caladea, serif', '"Segoe UI", Calibri, Carlito, Roboto, Arial, sans-serif', "Cambria", "Calibri"),
    "classic": ('Arial, "Liberation Sans", Helvetica, sans-serif', 'Arial, "Liberation Sans", Helvetica, sans-serif', "Arial", "Arial"),
}
e = lambda s: html.escape(str(s if s is not None else ""), quote=True)


def load(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def parse_date(s):
    try:
        return datetime.date.fromisoformat(str(s))
    except ValueError:
        return None


def nice_date(d):
    return "%d %s %d" % (d.day, d.strftime("%b"), d.year) if d else ""


# ---------- arithmetic ----------

def money(v, cur):
    return V.fmt(v, "currency", cur, compact=False)


def whole(x):
    """Round half up to a whole unit: proposals and invoices show whole rupees or dollars."""
    from decimal import Decimal, ROUND_HALF_UP
    return float(Decimal(str(x)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def compute(spec):
    """Every figure in the proposal, computed once, in whole currency units so
    every printed figure adds up exactly."""
    cur = spec.get("currency", "INR")
    pr = next((s for s in spec.get("sections", []) if s.get("type") == "pricing"), None)
    out = {"currency": cur, "groups": {}, "lines": [], "optional": []}
    if pr:
        for it in pr.get("items", []):
            qty = float(it.get("qty", 1))
            rate = float(it["rate"])
            amt = whole(qty * rate)
            row = dict(it, qty=qty, rate=rate, amount=amt, billing=it.get("billing", "one-time"))
            (out["optional"] if it.get("optional") else out["lines"]).append(row)
        tax = pr.get("tax", {"label": "GST", "rate": 0.18}) if cur == "INR" else pr.get("tax")
        for billing in dict.fromkeys(r["billing"] for r in out["lines"]):
            rows = [r for r in out["lines"] if r["billing"] == billing]
            sub = whole(sum(r["amount"] for r in rows))
            disc = 0.0
            d = (pr.get("discount") or {}) if billing == "one-time" else {}
            if d.get("percent"):
                disc = whole(sub * float(d["percent"]))
            elif d.get("amount"):
                disc = whole(float(d["amount"]))
            taxable = whole(sub - disc)
            t = whole(taxable * float(tax["rate"])) if tax else 0.0
            out["groups"][billing] = {"subtotal": sub, "discount": disc, "discount_label": d.get("label", "Discount"),
                                      "taxable": taxable, "tax": t, "tax_label": (tax or {}).get("label", ""),
                                      "tax_rate": (tax or {}).get("rate", 0), "total": whole(taxable + t)}
        one = out["groups"].get("one-time")
        sched = []
        for m in pr.get("payments", []):
            base = one["total"] if one else 0
            sched.append(dict(m, amount=whole(base * float(m["percent"]))))
        if sched and one:  # rounding goes on the last milestone so the schedule adds up exactly
            diff = one["total"] - sum(m["amount"] for m in sched)
            sched[-1]["amount"] = whole(sched[-1]["amount"] + diff)
        out["payments"] = sched
    tl = next((s for s in spec.get("sections", []) if s.get("type") == "timeline"), None)
    if tl and tl.get("phases"):
        ph = tl["phases"]
        if all("start" in p for p in ph):
            ds = [parse_date(p["start"]) for p in ph] + [parse_date(p["end"]) for p in ph]
            if all(ds):
                start, end = min(ds), max(ds)
                out["start"], out["end"] = start.isoformat(), end.isoformat()
                out["weeks"] = -(-((end - start).days + 1) // 7)  # whole weeks, rounded up
        elif all("weeks" in p for p in ph):
            out["weeks"] = max(p["weeks"][1] for p in ph)
    return out


def figures(nums):
    """Flat list of every computed number (for tracing text)."""
    vals = []
    for g in nums["groups"].values():
        vals += [g[k] for k in ("subtotal", "discount", "taxable", "tax", "total")] + [g["tax_rate"]]
    for r in nums["lines"] + nums["optional"]:
        vals += [r["qty"], r["rate"], r["amount"]]
    for m in nums.get("payments", []):
        vals += [m["amount"], float(m["percent"])]
    if "weeks" in nums:
        vals.append(nums["weeks"])
    return vals


# ---------- validation ----------

def validate(spec, bdir="."):
    p = []
    for k in ("title", "client", "date"):
        if not spec.get(k):
            p.append("%s is required" % k)
    if spec.get("date") and not parse_date(spec["date"]):
        p.append("date must be YYYY-MM-DD")
    if spec.get("valid_until"):
        vu, d = parse_date(spec["valid_until"]), parse_date(spec.get("date", ""))
        if not vu:
            p.append("valid_until must be YYYY-MM-DD")
        elif d and vu <= d:
            p.append("valid_until must be after the proposal date")
    if not (spec.get("from") or {}).get("name"):
        p.append("from.name (your company) is required")
    for key, allowed in (("theme", tuple(BK.CURATED)), ("motif", BK.MOTIFS), ("type", tuple(WEB_TYPES))):
        if spec.get(key) is not None and spec[key] not in allowed:
            p.append("%s must be one of %s" % (key, ", ".join(allowed)))
    b = spec.get("brand") or {}
    for key in ("primary", "accent"):
        if b.get(key):
            try:
                BK.hexc(b[key])
            except ValueError:
                p.append("brand %s must be a hex colour like #1F5FA8" % key)
    if b.get("logo"):
        pr = BK.logo_check(b["logo"] if os.path.isabs(b["logo"]) else os.path.join(bdir, b["logo"]))
        if pr:
            p.append(pr)
    secs = spec.get("sections", [])
    types = [s.get("type") for s in secs]
    for i, s in enumerate(secs, 1):
        t = s.get("type")
        tag = "section %d (%s)" % (i, t)
        if t not in SECTION_TYPES:
            p.append("%s: type must be one of %s" % (tag, ", ".join(SECTION_TYPES)))
            continue
        if not s.get("headline"):
            p.append("%s: needs a headline" % tag)
        if t == "pricing":
            items = s.get("items", [])
            if not items:
                p.append("%s: needs items (item, qty, rate)" % tag)
            for it in items:
                if not it.get("item") or not isinstance(it.get("rate"), (int, float)):
                    p.append("%s: every item needs a name and a numeric rate" % tag)
                if it.get("billing", "one-time") not in ("one-time", "monthly", "yearly"):
                    p.append("%s: billing must be one-time, monthly or yearly" % tag)
            pays = s.get("payments", [])
            if pays:
                tot = round(sum(float(m.get("percent", 0)) for m in pays), 6)
                if abs(tot - 1) > 1e-6:
                    p.append("%s: payment milestones add up to %s%%, not 100%%" % (tag, round(tot * 100, 2)))
            d = s.get("discount") or {}
            if d.get("percent") and not (0 < float(d["percent"]) < 1):
                p.append("%s: discount percent is a fraction, e.g. 0.1 for 10%%" % tag)
            if s.get("tax") and not (0 <= float(s["tax"].get("rate", -1)) < 1):
                p.append("%s: tax rate is a fraction, e.g. 0.18 for GST 18%%" % tag)
        if t == "timeline":
            ph = s.get("phases", [])
            if not ph:
                p.append("%s: needs phases" % tag)
            prev = None
            for q in ph:
                if "start" in q:
                    a, z = parse_date(q.get("start")), parse_date(q.get("end"))
                    if not a or not z:
                        p.append("%s: phase %r needs start and end as YYYY-MM-DD" % (tag, q.get("name")))
                    elif z < a:
                        p.append("%s: phase %r ends before it starts" % (tag, q.get("name")))
                    elif prev and a < prev:
                        p.append("%s: phase %r starts before the previous phase starts; list phases in order" % (tag, q.get("name")))
                    else:
                        prev = a
                        d = parse_date(spec.get("date", ""))
                        if d and a < d:
                            p.append("%s: phase %r starts before the proposal date" % (tag, q.get("name")))
                elif "weeks" in q:
                    w = q["weeks"]
                    if not (isinstance(w, list) and len(w) == 2 and 1 <= w[0] <= w[1]):
                        p.append("%s: phase %r weeks must be [first, last], e.g. [1, 3]" % (tag, q.get("name")))
                else:
                    p.append("%s: phase %r needs start/end dates or weeks" % (tag, q.get("name")))
            if ph and len({("start" in q) for q in ph}) > 1:
                p.append("%s: use dates for every phase or weeks for every phase, not both" % tag)
    if "pricing" not in types:
        p.append("a proposal needs a pricing section")
    if types and types[0] != "summary":
        p.append("the first section must be the summary (the answer first)")
    return p


# ---------- HTML ----------

CSS = """
@page { size: A4; margin: 22mm 20mm 24mm;
  @bottom-left { content: "%(foot)s"; font: 8.5pt %(fbody)s; color: #%(muted)s; }
  @bottom-right { content: "Page " counter(page) " of " counter(pages); font: 8.5pt %(fbody)s; color: #%(muted)s; } }
@page :first { margin: 0; @bottom-left { content: none; } @bottom-right { content: none; } }
* { box-sizing: border-box; }
html { -webkit-print-color-adjust: exact; print-color-adjust: exact; }
body { margin: 0; font: 10.5pt/1.55 %(fbody)s; color: #%(text)s; }
h1, h2, h3 { font-family: %(fhead)s; color: #%(text)s; line-height: 1.15; margin: 0; }
h2 { font-size: 19pt; margin: 0 0 10pt; letter-spacing: -0.01em; break-after: avoid; }
h3 { font-size: 12pt; margin: 14pt 0 4pt; break-after: avoid; }
p { margin: 0 0 8pt; }
.body { color: #%(body_c)s; }
section { margin-bottom: 22pt; }
section.newpage { break-before: page; }
ul { margin: 4pt 0 10pt; padding-left: 15pt; } li { margin: 3pt 0; } li::marker { color: #%(accent_fill)s; }
/* cover */
.cover { position: relative; width: 210mm; height: 296mm; background: #%(dark)s; color: #fff; overflow: hidden; break-after: page; }
.cover .orb { position: absolute; border-radius: 50%%; }
.cover .in { position: absolute; left: 20mm; right: 20mm; top: 22mm; bottom: 22mm; display: flex; flex-direction: column; }
.cover .logo { height: 13mm; align-self: flex-start; }
.cover .chip { background: #fff; border-radius: 8px; padding: 3mm 4mm; align-self: flex-start; }
.cover .chip img { height: 9mm; display: block; }
.cover .kicker { margin-top: auto; color: #%(accent)s; font-weight: 700; font-size: 11pt; }
.cover h1 { color: #fff; font-size: 32pt; margin-top: 5mm; max-width: 150mm; letter-spacing: -0.015em; }
.cover .for { margin-top: 7mm; font-size: 14pt; color: #%(on_dark2)s; }
.cover .meta { display: grid; grid-template-columns: repeat(3, auto); justify-content: start; gap: 2mm 12mm; margin-top: 16mm;
  font-size: 9.5pt; color: #%(on_dark2)s; padding-top: 6mm; border-top: 1px solid #%(dark_line)s; }
.cover .meta b { display: block; color: #fff; font-size: 10.5pt; }
/* summary */
.keys { display: grid; grid-template-columns: repeat(var(--n), 1fr); gap: 0; margin: 12pt 0 16pt; border-top: 2px solid #%(text)s; }
.keys div { padding: 9pt 10pt 0 0; }
.keys .v { font-family: %(fhead)s; font-size: 17pt; font-weight: 700; color: #%(accent_dark)s; line-height: 1.15; }
.keys .l { font-size: 9pt; color: #%(muted)s; margin-top: 2pt; }
.lede { font-size: 12pt; color: #%(text)s; }
/* tables */
table { width: 100%%; border-collapse: collapse; margin: 6pt 0 10pt; font-variant-numeric: tabular-nums; break-inside: auto; }
th { text-align: left; font-size: 8.5pt; font-weight: 700; color: #%(muted)s; padding: 5pt 6pt; border-bottom: 1.5px solid #%(text)s; }
td { padding: 6pt 6pt; border-bottom: 1px solid #%(grid)s; vertical-align: top; }
tr { break-inside: avoid; }
table.price, table.keep { break-inside: avoid; } /* totals stay with their lines; short tables stay whole */
h3 + table { break-before: avoid; }
td.n, th.n { text-align: right; white-space: nowrap; }
td .desc { display: block; color: #%(muted)s; font-size: 9pt; margin-top: 1pt; }
tr.sum td { border-bottom: 0; padding: 3pt 6pt; }
tr.sum td.lab { text-align: right; color: #%(body_c)s; }
tr.total td { font-weight: 700; font-size: 12pt; border-top: 1.5px solid #%(text)s; padding-top: 7pt; }
tr.total td.n { color: #%(accent_dark)s; }
.note { font-size: 9pt; color: #%(muted)s; }
/* scope */
.deliv { display: grid; grid-template-columns: 1fr; gap: 0; }
.deliv > div { padding: 8pt 0; border-top: 1px solid #%(grid)s; break-inside: avoid; }
.deliv b { display: block; }
.two { display: grid; grid-template-columns: 1fr 1fr; gap: 0 16pt; }
/* timeline */
.gantt { width: 100%%; margin: 6pt 0 8pt; }
/* team and proof */
.people { display: grid; grid-template-columns: 1fr 1fr; gap: 10pt 16pt; }
.people > div { break-inside: avoid; }
.people .role { color: #%(accent_dark)s; font-weight: 600; font-size: 9.5pt; }
.proof { display: grid; grid-template-columns: repeat(var(--n), 1fr); gap: 12pt; }
.proof > div { border-top: 2px solid #%(accent_fill)s; padding-top: 7pt; break-inside: avoid; }
.proof .r { font-family: %(fhead)s; font-weight: 700; font-size: 13pt; line-height: 1.2; }
.proof .c { color: #%(muted)s; font-size: 9pt; margin-top: 3pt; }
/* acceptance */
.sign { display: grid; grid-template-columns: 1fr 1fr; gap: 18pt; margin-top: 16pt; break-inside: avoid; }
.sign div { border-top: 1px solid #%(text)s; padding-top: 5pt; font-size: 9.5pt; color: #%(body_c)s; min-height: 30mm; }
.sign b { color: #%(text)s; }
"""


def tokens(spec, bdir):
    T, note = BK.resolve_spec(spec, bdir)
    head, body, fh, fb = WEB_TYPES.get(spec.get("type", "modern"), WEB_TYPES["modern"])
    return dict(T, fhead=head, fbody=body, font_head_docx=fh, font_body_docx=fb, body_c=T["body"],
                dark_line=BK.mix("FFFFFF", T["dark"], 0.2)), note


def paras(x):
    if not x:
        return ""
    xs = x if isinstance(x, list) else [x]
    return "".join("<p class='body'>%s</p>" % e(t) for t in xs)


def bullets(xs):
    return "<ul>%s</ul>" % "".join("<li>%s</li>" % e(x) for x in xs) if xs else ""


def summary_keys(spec, nums):
    cur = nums["currency"]
    keys = []
    one = nums["groups"].get("one-time")
    if one:
        keys.append((money(one["total"], cur), "Investment%s" % (" incl. %s" % one["tax_label"] if one["tax"] else "")))
    for billing, label in (("monthly", "per month"), ("yearly", "per year")):
        g = nums["groups"].get(billing)
        if g:
            keys.append((money(g["total"], cur), "%s%s" % (label.capitalize(), " incl. %s" % g["tax_label"] if g["tax"] else "")))
    if "weeks" in nums:
        w = nums["weeks"]
        keys.append(("%s weeks" % (int(w) if float(w).is_integer() else w), "Timeline"))
    if nums.get("start"):
        keys.append((nice_date(parse_date(nums["start"])), "Start"))
    return keys[:4]


def gantt_svg(s, T, nums):
    ph = s["phases"]
    W, rowh, lw = 640, 26, 170
    dated = all("start" in p for p in ph)
    if dated:
        a = min(parse_date(p["start"]) for p in ph)
        z = max(parse_date(p["end"]) for p in ph)
        span = (z - a).days + 1
        pos = lambda p: ((parse_date(p["start"]) - a).days, (parse_date(p["end"]) - a).days + 1)
    else:
        span = max(p["weeks"][1] for p in ph)
        pos = lambda p: (p["weeks"][0] - 1, p["weeks"][1])
    H = rowh * len(ph) + 26
    sx = lambda v: lw + (W - lw - 8) * v / span
    out = ['<svg class="gantt" viewBox="0 0 %d %d" role="img" aria-label="Project timeline">' % (W, H)]
    # week grid
    weeks = span / 7 if dated else span
    step = 1 if weeks <= 12 else (2 if weeks <= 24 else 4)
    k = 0
    while k <= weeks:
        x = sx(k * 7 if dated else k)
        out.append('<line x1="%.1f" x2="%.1f" y1="0" y2="%d" stroke="#%s" stroke-width="1"/>' % (x, x, H - 18, T["grid"]))
        if k < weeks:
            out.append('<text x="%.1f" y="%d" font-size="9" fill="#%s">W%d</text>' % (x + 3, H - 5, T["muted"], k + 1))
        k += step
    for i, p in enumerate(ph):
        y = i * rowh + 4
        s0, s1 = pos(p)
        out.append('<text x="0" y="%d" font-size="10.5" fill="#%s">%s</text>' % (y + 14, T["text"], e(p["name"])))
        out.append('<rect x="%.1f" y="%d" width="%.1f" height="16" rx="4" fill="#%s"/>' % (
            sx(s0), y + 2, max(4, sx(s1) - sx(s0)), T["accent_fill"] if i % 2 == 0 else T["dark"]))
    out.append("</svg>")
    return "".join(out)


def pricing_html(s, T, nums):
    cur = nums["currency"]
    out = []
    for billing, g in nums["groups"].items():
        rows = [r for r in nums["lines"] if r["billing"] == billing]
        title = {"one-time": "", "monthly": "Monthly", "yearly": "Yearly"}[billing]
        if title and len(nums["groups"]) > 1:
            out.append("<h3>%s</h3>" % title)
        body = "".join('<tr><td>%s%s</td><td class="n">%s</td><td class="n">%s</td><td class="n">%s</td></tr>' % (
            e(r["item"]), '<span class="desc">%s</span>' % e(r["description"]) if r.get("description") else "",
            e(V.fmt(r["qty"], "number")) + (" " + e(r["unit"]) if r.get("unit") else ""), e(money(r["rate"], cur)), e(money(r["amount"], cur))) for r in rows)
        sums = '<tr class="sum"><td colspan="3" class="lab">Subtotal</td><td class="n">%s</td></tr>' % e(money(g["subtotal"], cur))
        if g["discount"]:
            sums += '<tr class="sum"><td colspan="3" class="lab">%s</td><td class="n">&minus;%s</td></tr>' % (e(g["discount_label"]), e(money(g["discount"], cur)))
        if g["tax"]:
            sums += '<tr class="sum"><td colspan="3" class="lab">%s %s%%</td><td class="n">%s</td></tr>' % (
                e(g["tax_label"]), e(V._trim(g["tax_rate"] * 100)), e(money(g["tax"], cur)))
        sums += '<tr class="total"><td colspan="3">Total%s</td><td class="n">%s</td></tr>' % (
            {"one-time": "", "monthly": " per month", "yearly": " per year"}[billing], e(money(g["total"], cur)))
        out.append('<table class="price"><tr><th>Item</th><th class="n">Qty</th><th class="n">Rate</th><th class="n">Amount</th></tr>%s%s</table>' % (body, sums))
    if nums["optional"]:
        out.append("<h3>Optional add-ons</h3><table class=\"keep\"><tr><th>Item</th><th class=\"n\">Amount</th></tr>%s</table>" % "".join(
            '<tr><td>%s%s</td><td class="n">%s</td></tr>' % (e(r["item"]), '<span class="desc">%s</span>' % e(r["description"]) if r.get("description") else "",
                                                           e(money(r["amount"], cur) + ("" if r["billing"] == "one-time" else " / " + r["billing"].replace("ly", "")))) for r in nums["optional"]))
    if nums.get("payments"):
        out.append("<h3>Payment schedule</h3><table class=\"keep\"><tr><th>Milestone</th><th>When</th><th class=\"n\">Share</th><th class=\"n\">Amount</th></tr>%s</table>" % "".join(
            '<tr><td>%s</td><td>%s</td><td class="n">%s</td><td class="n">%s</td></tr>' % (
                e(m["milestone"]), e(m.get("due", "")), e(V.fmt(float(m["percent"]), "percent")), e(money(m["amount"], cur))) for m in nums["payments"]))
    if s.get("notes"):
        out.append('<p class="note">%s</p>' % "<br>".join(e(n) for n in (s["notes"] if isinstance(s["notes"], list) else [s["notes"]])))
    return "".join(out)


def section_html(s, spec, T, nums):
    t = s["type"]
    h = "<h2>%s</h2>" % e(s["headline"])
    if t == "summary":
        keys = summary_keys(spec, nums)
        kh = '<div class="keys" style="--n:%d">%s</div>' % (len(keys), "".join(
            '<div><div class="v">%s</div><div class="l">%s</div></div>' % (e(v), e(l)) for v, l in keys)) if keys else ""
        lede = '<p class="lede">%s</p>' % e(s["lede"]) if s.get("lede") else ""
        body = lede + kh + paras(s.get("body")) + bullets(s.get("points", []))
    elif t == "understanding":
        body = paras(s.get("body")) + ("<h3>Goals</h3>" + bullets(s["goals"]) if s.get("goals") else "") + \
            ("<h3>What stands in the way</h3>" + bullets(s["challenges"]) if s.get("challenges") else "")
    elif t == "scope":
        dl = "".join("<div><b>%s</b>%s</div>" % (e(d["name"]), "<span class='body'>%s</span>" % e(d.get("description", ""))) for d in s.get("deliverables", []))
        body = paras(s.get("body")) + ('<div class="deliv">%s</div>' % dl if dl else "")
        if s.get("included") or s.get("excluded"):
            body += '<div class="two"><div>%s</div><div>%s</div></div>' % (
                "<h3>Included</h3>" + bullets(s.get("included", [])) if s.get("included") else "",
                "<h3>Not included</h3>" + bullets(s.get("excluded", [])) if s.get("excluded") else "")
    elif t == "timeline":
        rows = "".join("<tr><td>%s</td><td>%s</td><td>%s</td></tr>" % (
            e(p["name"]), e("%s to %s" % (nice_date(parse_date(p["start"])), nice_date(parse_date(p["end"])))) if "start" in p else
            e("Week %d to %d" % tuple(p["weeks"]) if p["weeks"][0] != p["weeks"][1] else "Week %d" % p["weeks"][0]), e(p.get("outcome", ""))) for p in s["phases"])
        body = paras(s.get("body")) + gantt_svg(s, T, nums) + "<table><tr><th>Phase</th><th>When</th><th>Outcome</th></tr>%s</table>" % rows
    elif t == "pricing":
        body = paras(s.get("body")) + pricing_html(s, T, nums)
    elif t == "team":
        body = paras(s.get("body")) + '<div class="people">%s</div>' % "".join(
            "<div><b>%s</b><div class='role'>%s</div><p class='body'>%s</p></div>" % (e(x["name"]), e(x.get("role", "")), e(x.get("bio", ""))) for x in s.get("people", []))
    elif t == "proof":
        items = s.get("items", [])
        body = paras(s.get("body")) + '<div class="proof" style="--n:%d">%s</div>' % (min(3, len(items)) or 1, "".join(
            "<div><div class='r'>%s</div><div class='c'>%s</div></div>" % (e(x["result"]), e(x.get("client", ""))) for x in items))
    elif t == "terms":
        body = paras(s.get("body")) + "<ol>%s</ol>" % "".join("<li class='body'>%s</li>" % e(x) for x in s.get("items", []))
    elif t == "acceptance":
        fr = spec["from"]
        body = paras(s.get("body")) + '<div class="sign"><div><b>For %s</b><br>Name, signature and date</div><div><b>For %s</b><br>%s</div></div>' % (
            e(spec["client"]), e(fr["name"]), e(", ".join(x for x in (fr.get("signatory"), fr.get("role")) if x) or "Name, signature and date"))
    else:
        body = paras(s.get("body")) + bullets(s.get("points", []))
    return '<section class="%s">%s%s</section>' % ("newpage" if s.get("new_page") else "", h, body)


def cover_html(spec, T):
    fr = spec["from"]
    logo = ""
    if T.get("logo"):
        img = '<img src="%s" alt="%s">' % (BK.data_uri(T["logo"]), e(fr["name"]))
        logo = '<div class="chip">%s</div>' % img if BK.logo_hidden_share(T["logo"], T["dark"]) > 0.15 else img.replace("<img", '<img class="logo"')
    else:
        logo = '<div style="font-weight:700;font-size:14pt">%s</div>' % e(fr["name"])
    meta = [("Prepared for", spec["client"]), ("Prepared by", fr["name"]), ("Date", nice_date(parse_date(spec["date"])))]
    if spec.get("valid_until"):
        meta.append(("Valid until", nice_date(parse_date(spec["valid_until"]))))
    if spec.get("number"):
        meta.append(("Reference", spec["number"]))
    orb = '<div class="orb" style="width:190mm;height:190mm;right:-70mm;top:-60mm;background:#%s"></div>' % BK.mix("FFFFFF", T["dark"], 0.06)
    if T.get("motif") == "rings":
        orb = "".join('<div class="orb" style="width:%dmm;height:%dmm;right:%dmm;top:%dmm;border:0.5mm solid #%s"></div>' % (
            d, d, -40 - d // 4 + 40, -30 - d // 4 + 30, BK.mix("FFFFFF", T["dark"], 0.14)) for d in (200, 150, 100, 50))
    dot = '<div class="orb" style="width:11mm;height:11mm;left:20mm;top:150mm;background:#%s"></div>' % T["accent"]
    return '<div class="cover">%s%s<div class="in">%s<div class="kicker">%s</div><h1>%s</h1><div class="for">%s</div><div class="meta">%s</div></div></div>' % (
        orb, dot, logo, e(spec.get("kicker", "Proposal")), e(spec["title"]), e(spec.get("subtitle", "For " + spec["client"])),
        "".join("<div>%s<b>%s</b></div>" % (e(a), e(b)) for a, b in meta))


def build_html(spec, T, nums):
    fr = spec["from"]
    foot = "%s, proposal for %s" % (fr["name"], spec["client"])
    css = CSS % dict(T, foot=foot.replace('"', "'"))
    secs = "".join(section_html(s, spec, T, nums) for s in spec["sections"])
    return '<!doctype html><html lang="%s"><head><meta charset="utf-8"><title>%s</title><style>%s</style></head><body>%s<main>%s</main></body></html>' % (
        e(spec.get("language", "en")), e(spec["title"]), css, cover_html(spec, T), secs)


# ---------- Word ----------

def build_docx(spec, T, nums, path):
    try:
        from docx import Document
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt, RGBColor, Mm
    except ImportError:
        return False
    cur = nums["currency"]
    rgb = lambda h: RGBColor.from_string(h)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(20)
    sec.top_margin, sec.bottom_margin = Mm(22), Mm(22)
    sec.different_first_page_header_footer = True
    st = doc.styles["Normal"]
    st.font.name, st.font.size = T["font_body_docx"], Pt(10.5)
    st.font.color.rgb = rgb(T["text"])
    for name, size in (("Heading 1", 19), ("Heading 2", 12), ("Title", 30)):
        hs = doc.styles[name]
        hs.font.name, hs.font.size, hs.font.bold = T["font_head_docx"], Pt(size), True
        hs.font.color.rgb = rgb(T["text"])
        rpr = hs.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs"):
            fonts.set(qn(a), T["font_head_docx"])
        for a in ("w:asciiTheme", "w:hAnsiTheme"):
            if fonts.get(qn(a)) is not None:
                del fonts.attrib[qn(a)]

    def shade(cell, hexcolor):
        tcPr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:color"), "auto")
        sh.set(qn("w:fill"), hexcolor)
        tcPr.append(sh)

    def para(text, size=None, color=None, bold=False, after=6, align=None):
        p = doc.add_paragraph()
        r = p.add_run(str(text))
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = rgb(color)
        r.bold = bold
        p.paragraph_format.space_after = Pt(after)
        if align:
            p.alignment = align
        return p

    def bl(items, style="List Bullet"):
        for x in items:
            doc.add_paragraph(str(x), style=style)

    def table(headers, rows, num_cols=(), widths=None):
        t = doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.style = "Table Grid"
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]
            c.text = ""
            r = c.paragraphs[0].add_run(h)
            r.bold, r.font.size = True, Pt(9)
            r.font.color.rgb = rgb("FFFFFF")
            shade(c, T["dark"])
            if i in num_cols:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                txt, b = (v if isinstance(v, tuple) else (v, False))
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(str(txt))
                r.font.size, r.bold = Pt(10), b
                if i in num_cols:
                    cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if widths:
            t.autofit = False
            grid = t._tbl.tblGrid
            for i, w in enumerate(widths):
                t.columns[i].width = Mm(w)
                if i < len(grid.gridCol_lst):
                    grid.gridCol_lst[i].w = Mm(w)
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Mm(w)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    # cover
    fr = spec["from"]
    if T.get("logo"):
        try:
            doc.add_picture(T["logo"], height=Mm(14))
        except Exception:
            pass
    for _ in range(6):
        doc.add_paragraph()
    para(spec.get("kicker", "Proposal"), 12, T["accent_dark"], True, 4)
    doc.add_paragraph(spec["title"], style="Title")
    para(spec.get("subtitle", "For " + spec["client"]), 14, T["body"], after=30)
    meta = [("Prepared for", spec["client"]), ("Prepared by", fr["name"]), ("Date", nice_date(parse_date(spec["date"])))]
    if spec.get("valid_until"):
        meta.append(("Valid until", nice_date(parse_date(spec["valid_until"]))))
    if spec.get("number"):
        meta.append(("Reference", spec["number"]))
    for a, b in meta:
        p = doc.add_paragraph()
        r1 = p.add_run(a + ": ")
        r1.font.color.rgb = rgb(T["muted"])
        p.add_run(b).bold = True
        p.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    # footer with page number on later pages
    fp = sec.footer.paragraphs[0]
    fp.text = "%s, proposal for %s    Page " % (fr["name"], spec["client"])
    fp.runs[0].font.size = Pt(8.5)
    fp.runs[0].font.color.rgb = rgb(T["muted"])
    for kind in ("begin", None, "end"):
        r = fp.add_run()
        r.font.size = Pt(8.5)
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = "PAGE"
            r._r.append(it)
    for s in spec["sections"]:
        t = s["type"]
        if s.get("new_page"):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading(s["headline"], level=1)
        if t == "summary":
            if s.get("lede"):
                para(s["lede"], 12)
            keys = summary_keys(spec, nums)
            if keys:
                table([l for _, l in keys], [[(v, True) for v, _ in keys]], widths=[170 / len(keys)] * len(keys))
        if s.get("body"):
            for x in (s["body"] if isinstance(s["body"], list) else [s["body"]]):
                para(x, color=T["body"])
        if t in ("summary", "text") and s.get("points"):
            bl(s["points"])
        if t == "understanding":
            if s.get("goals"):
                doc.add_heading("Goals", level=2)
                bl(s["goals"])
            if s.get("challenges"):
                doc.add_heading("What stands in the way", level=2)
                bl(s["challenges"])
        if t == "scope":
            for d in s.get("deliverables", []):
                p = doc.add_paragraph()
                p.add_run(d["name"]).bold = True
                if d.get("description"):
                    p.add_run("\n" + d["description"]).font.color.rgb = rgb(T["body"])
            if s.get("included"):
                doc.add_heading("Included", level=2)
                bl(s["included"])
            if s.get("excluded"):
                doc.add_heading("Not included", level=2)
                bl(s["excluded"])
        if t == "timeline":
            table(["Phase", "When", "Outcome"], [[p["name"], ("%s to %s" % (nice_date(parse_date(p["start"])), nice_date(parse_date(p["end"])))) if "start" in p
                                                  else ("Week %d to %d" % tuple(p["weeks"]) if p["weeks"][0] != p["weeks"][1] else "Week %d" % p["weeks"][0]), p.get("outcome", "")] for p in s["phases"]],
                  widths=(55, 50, 65))
        if t == "pricing":
            for billing, g in nums["groups"].items():
                if len(nums["groups"]) > 1:
                    doc.add_heading({"one-time": "One-time", "monthly": "Monthly", "yearly": "Yearly"}[billing], level=2)
                rows = [[r["item"] + ("\n" + r["description"] if r.get("description") else ""), V.fmt(r["qty"], "number") + (" " + r["unit"] if r.get("unit") else ""),
                         money(r["rate"], cur), money(r["amount"], cur)] for r in nums["lines"] if r["billing"] == billing]
                rows.append(["", "", "Subtotal", money(g["subtotal"], cur)])
                if g["discount"]:
                    rows.append(["", "", g["discount_label"], "-" + money(g["discount"], cur)])
                if g["tax"]:
                    rows.append(["", "", "%s %s%%" % (g["tax_label"], V._trim(g["tax_rate"] * 100)), money(g["tax"], cur)])
                rows.append([("Total", True), "", "", (money(g["total"], cur), True)])
                table(["Item", "Qty", "Rate", "Amount"], rows, num_cols=(1, 2, 3), widths=(80, 25, 32, 33))
            if nums["optional"]:
                doc.add_heading("Optional add-ons", level=2)
                table(["Item", "Amount"], [[r["item"], money(r["amount"], cur)] for r in nums["optional"]], num_cols=(1,), widths=(130, 40))
            if nums.get("payments"):
                doc.add_heading("Payment schedule", level=2)
                table(["Milestone", "When", "Share", "Amount"], [[m["milestone"], m.get("due", ""), V.fmt(float(m["percent"]), "percent"), money(m["amount"], cur)]
                                                                 for m in nums["payments"]], num_cols=(2, 3), widths=(60, 50, 25, 35))
            for n in (s.get("notes") if isinstance(s.get("notes"), list) else [s["notes"]] if s.get("notes") else []):
                para(n, 9, T["muted"])
        if t == "team":
            for x in s.get("people", []):
                p = doc.add_paragraph()
                p.add_run(x["name"]).bold = True
                if x.get("role"):
                    rr = p.add_run("  " + x["role"])
                    rr.font.color.rgb = rgb(T["accent_dark"])
                if x.get("bio"):
                    p.add_run("\n" + x["bio"])
        if t == "proof":
            for x in s.get("items", []):
                p = doc.add_paragraph()
                p.add_run(x["result"]).bold = True
                if x.get("client"):
                    p.add_run("\n" + x["client"]).font.color.rgb = rgb(T["muted"])
        if t == "terms":
            bl(s.get("items", []), "List Number")
        if t == "acceptance":
            for _ in range(2):
                doc.add_paragraph()
            tb = doc.add_table(rows=1, cols=2)
            tb.rows[0].cells[0].text = "For %s\n\n\n____________________________\nName, signature and date" % spec["client"]
            tb.rows[0].cells[1].text = "For %s\n\n\n____________________________\n%s" % (fr["name"], ", ".join(x for x in (fr.get("signatory"), fr.get("role")) if x) or "Name, signature and date")
    doc.core_properties.title = spec["title"]
    doc.core_properties.author = fr["name"]
    doc.save(path)
    return True


def main():
    a = sys.argv[1:]
    if not a:
        sys.exit(__doc__)
    src = a[0]
    out = a[a.index("--out") + 1] if "--out" in a else "proposal"
    spec = load(src)
    bdir = os.path.dirname(os.path.abspath(src))
    probs = validate(spec, bdir)
    if probs:
        sys.exit("spec problems:\n- " + "\n- ".join(probs))
    T, note = tokens(spec, bdir)
    nums = compute(spec)
    with open(out + ".numbers.json", "w", encoding="utf-8") as f:
        json.dump(nums, f, indent=1, ensure_ascii=False)
    with open(out + ".html", "w", encoding="utf-8") as f:
        f.write(build_html(spec, T, nums))
    made = [out + ".html"]
    if "--no-pdf" not in a:
        try:
            render.html_to_pdf(out + ".html", out + ".pdf")
            made.append(out + ".pdf")
        except Exception as ex:  # noqa: BLE001
            print("PDF skipped: %s" % ex)
    if "--no-docx" not in a:
        if build_docx(spec, T, nums, out + ".docx"):
            made.append(out + ".docx")
        else:
            print("Word file skipped: pip install python-docx")
    if note:
        print(note + " (tell the user; they can set brand.primary and brand.accent to change them)")
    one = nums["groups"].get("one-time")
    if one:
        print("total %s (subtotal %s, %s %s)" % (money(one["total"], nums["currency"]), money(one["subtotal"], nums["currency"]),
                                                  one["tax_label"] or "no tax", money(one["tax"], nums["currency"]) if one["tax"] else ""))
    print("built " + ", ".join(made))


if __name__ == "__main__":
    main()
